import csv
import re
from pathlib import Path

import pymupdf
import pytest
from openpyxl import load_workbook

from anonimizador.lote import Lote, desde_utf16, mapa_utf16
from anonimizador.memoria import fragmentar

TEXTO = (
    "HISTORIA CLÍNICA\nPaciente: Laura Martínez Ospina\nCC 1067618172\n"
    "Teléfono: 3001234567\nDiagnóstico: cefalea tensional."
)


def crear_pdf(ruta: Path, *paginas: str) -> Path:
    doc = pymupdf.open()
    for texto in paginas or (TEXTO,):
        doc.new_page().insert_text((72, 72), texto)
    doc.save(ruta)
    return ruta


def texto_pdf(ruta: Path) -> str:
    return "\n".join(p.get_text() for p in pymupdf.open(ruta))


def test_lote_completo(tmp_path, anonimizador):
    origen = tmp_path / "origen"
    origen.mkdir()
    pdfs = [crear_pdf(origen / f"Historia {n} - Martínez.pdf") for n in (1, 2)]
    (origen / "vacio.pdf").write_bytes(b"no es un pdf")
    pdfs.append(origen / "vacio.pdf")

    eventos = []
    lote = Lote.crear(pdfs, tmp_path / "salida")
    lote.procesar(anonimizador, al_avanzar=lambda e: eventos.append((e.codigo, e.estado)))

    c = lote.almacen.contadores()
    assert (c["total"], c["completados"], c["errores"], c["pendientes"]) == (3, 2, 1, 0)

    total, docs = lote.almacen.listar(0, 10)
    codigos = [d["codigo"] for d in docs]
    assert total == 3 and len(set(codigos)) == 3
    assert all(re.fullmatch(r"HC-[2-9A-HJKMNP-Z]{8}", c) for c in codigos)
    for d in docs:
        if d["estado"] != "completado":
            assert d["error"]
            continue
        pdf = lote.carpeta_pdfs / f"{d['codigo']}.pdf"
        texto = texto_pdf(pdf)
        assert "Laura" not in texto and "1067618172" not in texto and "3001234567" not in texto
        assert "cefalea" in texto
        assert "Martínez" not in str(pymupdf.open(pdf).metadata)

    hoja = load_workbook(lote.carpeta / "relacion_codigos.xlsx").active
    filas = list(hoja.iter_rows(min_row=2, values_only=True))
    assert {(f[0], f[1]) for f in filas} == {(d["codigo"], d["nombre_original"]) for d in docs}
    with open(lote.carpeta / "relacion_codigos.csv", encoding="utf-8-sig") as f:
        assert len(list(csv.reader(f, delimiter=";"))) == 4
    assert [e for _, e in eventos].count("procesando") == 3 and len(eventos) == 6


def test_pausar_y_reanudar(tmp_path, anonimizador):
    """Pausar deja el resto pendiente; reabrir el lote (p. ej. tras cerrar la
    app) continúa donde iba, sin reprocesar lo hecho."""
    pdfs = [crear_pdf(tmp_path / f"{n}.pdf") for n in range(4)]
    lote = Lote.crear(pdfs, tmp_path / "salida")
    hechos = []
    lote.procesar(anonimizador, al_avanzar=lambda e: hechos.append(e.estado),
                  cancelado=lambda: hechos.count("completado") >= 1)
    assert lote.almacen.contadores()["pendientes"] == 3
    filas = list(load_workbook(lote.carpeta / "relacion_codigos.xlsx").active.values)[1:]
    assert [f[4] for f in filas] == ["Anonimizado", "Pendiente", "Pendiente", "Pendiente"]

    # Simula un cierre abrupto con un documento a mitad de proceso.
    siguiente = lote.almacen.siguiente_pendiente()
    lote.almacen.marcar_procesando(siguiente["id"])

    reabierto = Lote.abrir(lote.carpeta)
    procesados = []
    reabierto.procesar(anonimizador, al_avanzar=lambda e: procesados.append(e.codigo))
    assert reabierto.almacen.contadores()["completados"] == 4
    assert len(set(procesados)) == 3  # solo los 3 que faltaban


def test_paginas_se_conservan(tmp_path, anonimizador):
    pdf = crear_pdf(tmp_path / "tres.pdf", "Pagina uno de Laura Gómez", "Pagina dos", "Pagina tres")
    lote = Lote.crear([pdf], tmp_path / "salida")
    lote.procesar(anonimizador)
    _, (doc,) = lote.almacen.listar(0, 1)
    assert doc["paginas"] == 3
    assert pymupdf.open(lote.carpeta_pdfs / f"{doc['codigo']}.pdf").page_count == 3


def test_carpetas_de_lote_no_se_pisan(tmp_path):
    a = Lote.crear([], tmp_path)
    b = Lote.crear([], tmp_path)
    assert a.carpeta != b.carpeta


# --------------------------------------------------------------------------- #
# Revisión
# --------------------------------------------------------------------------- #

@pytest.fixture
def lote_revision(tmp_path, anonimizador):
    pdf = crear_pdf(
        tmp_path / "hc.pdf",
        "Paciente: Laura Martínez Ospina en cama 12.",
        "Trasladar desde cama 12 a piso 3. Diagnóstico: cefalea.",
    )
    lote = Lote.crear([pdf], tmp_path / "salida")
    lote.procesar(anonimizador)
    codigo = lote.almacen.listar(0, 1)[1][0]["codigo"]
    return lote, codigo


def test_ocultar_seleccion_en_todas_las_apariciones(lote_revision):
    lote, codigo = lote_revision
    doc = lote.documento_revision(codigo)
    assert "cama 12" in texto_pdf(lote.carpeta_pdfs / f"{codigo}.pdf")  # el detector no lo ve

    inicio = doc["texto"].index("cama 12")
    # La selección trae el espacio previo: se recorta.
    assert lote.ocultar(codigo, inicio - 1, inicio + len("cama 12"), "direcciones", todas=True) == 2
    salida = texto_pdf(lote.carpeta_pdfs / f"{codigo}.pdf")
    assert "cama 12" not in salida and salida.count("[DIRECCION]") == 2
    assert pymupdf.open(lote.carpeta_pdfs / f"{codigo}.pdf").page_count == 2
    manuales = [t for t in lote.documento_revision(codigo)["tramos"] if t["origen"] == "manual"]
    assert len(manuales) == 2


def test_ocultar_absorbe_tramos_solapados(lote_revision):
    lote, codigo = lote_revision
    doc = lote.documento_revision(codigo)
    inicio = doc["texto"].index("Paciente")
    fin = doc["texto"].index("Ospina") + len("Ospina")
    lote.ocultar(codigo, inicio, fin, "nombres", todas=False)
    tramos = lote.documento_revision(codigo)["tramos"]
    cubren = [t for t in tramos if t["inicio"] < fin and inicio < t["fin"]]
    assert len(cubren) == 1 and (cubren[0]["inicio"], cubren[0]["fin"]) == (inicio, fin)


def test_revertir_y_cambiar_categoria(lote_revision):
    lote, codigo = lote_revision
    nombre = next(t for t in lote.documento_revision(codigo)["tramos"] if t["categoria"] == "nombres")
    lote.actualizar_tramo(codigo, nombre["id"], activo=False)
    assert "Laura Martínez Ospina" in texto_pdf(lote.carpeta_pdfs / f"{codigo}.pdf")

    lote.actualizar_tramo(codigo, nombre["id"], activo=True, categoria="ubicaciones")
    salida = texto_pdf(lote.carpeta_pdfs / f"{codigo}.pdf")
    assert "Laura" not in salida and "[UBICACION]" in salida

    lote.marcar_revisado(codigo, True)
    assert lote.almacen.contadores()["revisados"] == 1


def test_offsets_utf16():
    texto = "a😀b Laura"
    a_js = mapa_utf16(texto)
    py = texto.index("Laura")
    assert a_js(py) == py + 1  # el emoji ocupa 2 unidades UTF-16
    assert desde_utf16(texto, a_js(py)) == py
    assert mapa_utf16("sin emojis")(4) == 4


def test_fragmentar_corta_en_saltos_de_linea():
    texto = "\n".join(f"linea {i} con texto" for i in range(2000))
    partes = fragmentar(texto, 1000)
    assert "".join(p for _, p in partes) == texto
    assert all(len(p) <= 1000 for _, p in partes)
    assert all(p.endswith("\n") for _, p in partes[:-1])
    assert all(texto[o:o + len(p)] == p for o, p in partes)
